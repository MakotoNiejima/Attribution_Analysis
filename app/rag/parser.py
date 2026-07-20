"""文档解析器：从 PDF / DOCX / XLSX / TXT 中提取纯文本。"""

from pathlib import Path


def parse_file(file_path: str | Path) -> str:
    """根据文件扩展名分发到对应的解析函数，返回纯文本。"""
    path = Path(file_path)
    suffix = path.suffix.lower()

    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".xls": _parse_xlsx,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".csv": _parse_text,
    }

    parser = parsers.get(suffix)
    if parser is None:
        raise ValueError(f"不支持的文件格式: {suffix}（支持 PDF/DOCX/XLSX/TXT/MD/CSV）")

    text = parser(path)
    # 去除多余空行，保留段落结构
    lines = text.splitlines()
    cleaned = []
    prev_blank = False
    for line in lines:
        is_blank = not line.strip()
        if is_blank and prev_blank:
            continue
        cleaned.append(line)
        prev_blank = is_blank
    return "\n".join(cleaned).strip()


def _parse_pdf(path: Path) -> str:
    """解析 PDF，逐页提取文本。"""
    import pypdf

    reader = pypdf.PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[第{i+1}页]\n{text}")
    return "\n\n".join(pages)


def _parse_docx(path: Path) -> str:
    """解析 Word 文档，按段落提取。"""
    from docx import Document

    doc = Document(str(path))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return "\n\n".join(paragraphs)


def _parse_xlsx(path: Path) -> str:
    """解析 Excel，逐行读取所有工作表。"""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    sheets_text = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join(cells)
            if line.strip(" |"):
                rows.append(line)
        if rows:
            sheets_text.append(f"[工作表: {sheet_name}]\n" + "\n".join(rows))

    wb.close()
    return "\n\n".join(sheets_text)


def _parse_text(path: Path) -> str:
    """直接读取纯文本文件。"""
    # 尝试 UTF-8，回退 GBK
    for encoding in ("utf-8", "gbk", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError(f"无法识别文件编码: {path}")
